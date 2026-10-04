from io import BytesIO
from pathlib import PurePosixPath
from urllib.parse import urlparse
from bs4 import BeautifulSoup
from zipfile import ZipFile

def check_archive(body):
    with ZipFile(BytesIO(body)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 64*1024*1024:
            raise ValueError("Office archive exceeds expanded size limit")

def parse_document(body,url,mime):
    suffix=PurePosixPath(urlparse(url).path).suffix.lower()
    if 'pdf' in mime or suffix=='.pdf':
        import fitz
        with fitz.open(stream=body,filetype='pdf') as doc:
            if len(doc)>100: raise ValueError('PDF exceeds 100 pages')
            return [{'page':i+1,'text':page.get_text()} for i,page in enumerate(doc)]
    if suffix=='.docx' or 'wordprocessingml' in mime:
        check_archive(body)
        from docx import Document
        doc=Document(BytesIO(body))
        text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
        return [{'page':None,'text':text}]
    if suffix=='.xlsx' or 'spreadsheetml' in mime:
        check_archive(body)
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(body),read_only=True,data_only=True); pages=[]
        try:
            for sheet in book:
                rows=[]
                for idx,row in enumerate(sheet.iter_rows(values_only=True)):
                    if idx>=10000: raise ValueError('Spreadsheet exceeds row limit')
                    rows.append(' | '.join('' if c is None else str(c) for c in row))
                pages.append({'page':None,'sheet':sheet.title,'text':'\n'.join(rows)})
        finally: book.close()
        return pages
    if 'html' in mime or suffix in ('','.html','.htm'):
        soup=BeautifulSoup(body,'html.parser')
        for n in soup(['script','style','nav','footer']): n.decompose()
        return [{'page':None,'text':soup.get_text('\n',strip=True)}]
    raise ValueError('Unsupported document type')
